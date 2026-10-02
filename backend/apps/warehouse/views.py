"""
仓库管理视图
"""
import hashlib
import json
import logging
import io
import uuid
from decimal import Decimal
from django.db import transaction
from django.http import HttpResponse
from django.utils import timezone
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from apps.core.response import success_response, error_response
from .models import (
    Unit, Category, Variety, Goods, StockIn, StockOut, Warning, Approval,
    Investigation, DestructionPlan, DestructionPlanItem,
    DestructionStageSnapshot, DestructionCorrection,
)
from .serializers import (
    UnitSerializer, UnitCreateSerializer,
    CategorySerializer, CategoryCreateSerializer,
    VarietySerializer, VarietyCreateSerializer,
    GoodsSerializer, StockInSerializer, StockOutSerializer,
    WarningSerializer, ApprovalSerializer,
    InvestigationSerializer, InvestigationCreateSerializer,
    DestructionPlanSerializer, DestructionPlanCreateSerializer,
    DestructionSnapshotSerializer,
    DestructionCorrectionSerializer, DestructionCorrectionCreateSerializer,
)

logger = logging.getLogger('apps')


def _first_error(errors):
    """从序列化器错误中提取第一条错误信息"""
    if isinstance(errors, dict):
        for value in errors.values():
            message = _first_error(value)
            if message:
                return message
    elif isinstance(errors, list):
        for value in errors:
            message = _first_error(value)
            if message:
                return message
    else:
        return str(errors)
    return ''


# ==================== 单位管理 ====================

class UnitListView(APIView):
    """单位列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Unit.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        units = queryset[start:end]
        
        serializer = UnitSerializer(units, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建单位"""
        serializer = UnitCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.create(
            name=serializer.validated_data['name'],
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='创建成功')


class UnitDetailView(APIView):
    """单位详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        serializer = UnitCreateSerializer(data=request.data, context={'instance': unit})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit.name = serializer.validated_data['name']
        unit.save()
        
        logger.info(f"User {request.user.username} updated unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        if unit.is_linked:
            return error_response(message='该单位已被关联，无法删除')
        
        name = unit.name
        unit.delete()
        
        logger.info(f"User {request.user.username} deleted unit {name}")
        
        return success_response(message='删除成功')


class UnitBatchDeleteView(APIView):
    """单位批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的单位')
        
        # 只删除未关联的单位
        units = Unit.objects.filter(pk__in=ids)
        deleted_count = 0
        for unit in units:
            if not unit.is_linked:
                unit.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} units")
        
        return success_response(message=f'成功删除 {deleted_count} 个单位')


class UnitAllView(APIView):
    """获取所有单位（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        units = Unit.objects.filter(is_active=True).order_by('name')
        serializer = UnitSerializer(units, many=True)
        return success_response(data=serializer.data)


# ==================== 品类管理 ====================

class CategoryListView(APIView):
    """品类列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Category.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        categories = queryset[start:end]
        
        serializer = CategorySerializer(categories, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品类"""
        serializer = CategoryCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category = Category.objects.create(
            name=serializer.validated_data['name'],
            unit=unit,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='创建成功')


class CategoryDetailView(APIView):
    """品类详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        serializer = CategoryCreateSerializer(data=request.data, context={'instance': category})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        category.name = serializer.validated_data['name']
        category.unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category.save()
        
        logger.info(f"User {request.user.username} updated category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        if category.is_linked:
            return error_response(message='该品类已被关联，无法删除')
        
        name = category.name
        category.delete()
        
        logger.info(f"User {request.user.username} deleted category {name}")
        
        return success_response(message='删除成功')


class CategoryBatchDeleteView(APIView):
    """品类批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品类')
        
        categories = Category.objects.filter(pk__in=ids)
        deleted_count = 0
        for category in categories:
            if not category.is_linked:
                category.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} categories")
        
        return success_response(message=f'成功删除 {deleted_count} 个品类')


class CategoryAllView(APIView):
    """获取所有品类（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        categories = Category.objects.filter(is_active=True).order_by('name')
        serializer = CategorySerializer(categories, many=True)
        return success_response(data=serializer.data)


# ==================== 品种管理 ====================

class VarietyListView(APIView):
    """品种列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Variety.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        varieties = queryset[start:end]
        
        serializer = VarietySerializer(varieties, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品种"""
        serializer = VarietyCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        category = Category.objects.get(pk=serializer.validated_data['category'])
        variety = Variety.objects.create(
            name=serializer.validated_data['name'],
            category=category,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='创建成功')


class VarietyDetailView(APIView):
    """品种详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        serializer = VarietyCreateSerializer(data=request.data, context={'instance': variety})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        variety.name = serializer.validated_data['name']
        variety.category = Category.objects.get(pk=serializer.validated_data['category'])
        variety.save()
        
        logger.info(f"User {request.user.username} updated variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        if variety.is_in_stock:
            return error_response(message='该品种已入库，无法删除')
        
        name = variety.name
        variety.delete()
        
        logger.info(f"User {request.user.username} deleted variety {name}")
        
        return success_response(message='删除成功')


class VarietyBatchDeleteView(APIView):
    """品种批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品种')
        
        varieties = Variety.objects.filter(pk__in=ids)
        deleted_count = 0
        for variety in varieties:
            if not variety.is_in_stock:
                variety.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} varieties")
        
        return success_response(message=f'成功删除 {deleted_count} 个品种')


class VarietyTemplateView(APIView):
    """品种导入模板下载"""
    permission_classes = []  # 允许匿名访问，通过token参数验证
    
    def get(self, request):
        # 从URL参数获取token进行验证
        from apps.authentication.backends import decode_token
        from apps.authentication.models import User
        
        token = request.query_params.get('token')
        if not token:
            return error_response(message='缺少认证信息', code=401)
        
        payload = decode_token(token)
        if not payload:
            return error_response(message='认证信息无效或已过期', code=401)
        
        try:
            user = User.objects.get(pk=payload['user_id'])
        except User.DoesNotExist:
            return error_response(message='用户不存在', code=401)
        
        wb = Workbook()
        
        # 第一个表格 - 导入模板
        ws1 = wb.active
        ws1.title = '品种导入'
        
        # 设置表头样式
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = ['品种', '品类', '单位']
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 设置列宽
        ws1.column_dimensions['A'].width = 25
        ws1.column_dimensions['B'].width = 20
        ws1.column_dimensions['C'].width = 15
        
        # 第二个表格 - 品类参考
        ws2 = wb.create_sheet(title='品类参考')
        
        headers2 = ['品类', '单位']
        for col, header in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 填充品类数据
        categories = Category.objects.filter(is_active=True).select_related('unit')
        for row, category in enumerate(categories, 2):
            ws2.cell(row=row, column=1, value=category.name).border = thin_border
            ws2.cell(row=row, column=2, value=category.unit.name).border = thin_border
        
        ws2.column_dimensions['A'].width = 20
        ws2.column_dimensions['B'].width = 15
        
        # 返回Excel文件
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=variety_import_template.xlsx'
        
        return response


class VarietyImportView(APIView):
    """品种导入视图"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    
    def post(self, request):
        if 'file' not in request.FILES:
            return error_response(message='请上传文件')
        
        file = request.FILES['file']
        
        try:
            wb = load_workbook(file)
            ws = wb.active
        except Exception as e:
            return error_response(message='文件格式错误，请上传Excel文件')
        
        # 获取所有品类及其单位
        categories = {c.name: c for c in Category.objects.filter(is_active=True).select_related('unit')}
        
        can_import = []
        cannot_import = []
        
        for row in range(2, ws.max_row + 1):
            variety_name = ws.cell(row=row, column=1).value
            category_name = ws.cell(row=row, column=2).value
            unit_name = ws.cell(row=row, column=3).value
            
            if not variety_name:
                continue
            
            variety_name = str(variety_name).strip()
            category_name = str(category_name).strip() if category_name else ''
            unit_name = str(unit_name).strip() if unit_name else ''
            
            # 验证
            error_msg = None
            
            if not variety_name:
                error_msg = '品种名称不能为空'
            elif len(variety_name) > 20:
                error_msg = '品种名称最多20个字'
            elif not category_name:
                error_msg = '品类不能为空'
            elif category_name not in categories:
                error_msg = f'品类"{category_name}"不存在'
            elif not unit_name:
                error_msg = '单位不能为空'
            elif categories.get(category_name) and categories[category_name].unit.name != unit_name:
                error_msg = f'单位与品类不匹配，应为"{categories[category_name].unit.name}"'
            elif Variety.objects.filter(name=variety_name, category__name=category_name).exists():
                error_msg = '该品种已存在'
            
            if error_msg:
                cannot_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name,
                    'reason': error_msg
                })
            else:
                can_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name
                })
        
        # 如果是预览请求
        if request.data.get('preview') == 'true':
            return success_response(data={
                'can_import': can_import,
                'cannot_import': cannot_import,
                'can_import_count': len(can_import),
                'cannot_import_count': len(cannot_import)
            })
        
        # 执行导入
        imported_count = 0
        for item in can_import:
            category = categories[item['category']]
            Variety.objects.create(
                name=item['variety'],
                category=category,
                created_by=request.user
            )
            imported_count += 1
        
        logger.info(f"User {request.user.username} imported {imported_count} varieties")
        
        return success_response(
            data={
                'imported_count': imported_count,
                'failed_count': len(cannot_import),
                'failed_items': cannot_import
            },
            message=f'成功导入 {imported_count} 个品种'
        )


# ==================== 其他视图占位 ====================

class DashboardView(APIView):
    """仪表盘视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'message': '仪表盘功能开发中...'
        })


class GoodsListView(APIView):
    """货物列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class StockInListView(APIView):
    """入库记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class StockOutListView(APIView):
    """出库记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class WarningListView(APIView):
    """预警记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class ApprovalListView(APIView):
    """审批记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


# ==================== 销毁管理 ====================

def _destruction_blockers(goods, item):
    """检查物资当前不满足销毁条件的原因，返回空列表表示可以销毁"""
    problems = []
    today = timezone.now().date()
    if goods.custody_until is None:
        problems.append('保管期限未设定')
    elif goods.custody_until > today:
        problems.append('保管期限未届满')
    if goods.is_frozen:
        problems.append('物资处于冻结状态')
    if goods.investigations.filter(status='open').exists():
        problems.append('存在未结调查')
    if goods.stock_outs.filter(status__in=['pending', 'approved']).exists():
        problems.append('存在未完成的领用记录')
    if item.quantity > goods.quantity:
        problems.append('库存数量不足')
    return problems


def _save_plan_snapshot(plan, stage, user):
    """保存指定阶段当时的清单摘要及摘要值"""
    items = list(plan.items.select_related('goods').order_by('id'))
    item_summaries = [
        {
            'item_id': item.id,
            'goods_code': item.goods.code,
            'goods_name': item.goods.name,
            'quantity': str(item.quantity),
            'status': item.status,
        }
        for item in items
    ]
    summary = {
        'plan_no': plan.plan_no,
        'stage': stage,
        'plan_status': plan.status,
        'item_count': len(item_summaries),
        'total_quantity': str(sum((item.quantity for item in items), Decimal('0'))),
        'items': item_summaries,
    }
    digest = hashlib.sha256(
        json.dumps(summary, ensure_ascii=False, sort_keys=True).encode('utf-8')
    ).hexdigest()
    return DestructionStageSnapshot.objects.create(
        plan=plan, stage=stage, summary=summary, digest=digest, created_by=user
    )


class GoodsFreezeView(APIView):
    """货物冻结/解冻视图"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            goods = Goods.objects.get(pk=pk)
        except Goods.DoesNotExist:
            return error_response(message='货物不存在', code=404)

        frozen = request.data.get('frozen')
        if not isinstance(frozen, bool):
            return error_response(message='请指定冻结状态（frozen: true/false）')
        reason = str(request.data.get('reason', '')).strip()
        if frozen and not reason:
            return error_response(message='请填写冻结原因')

        goods.is_frozen = frozen
        goods.freeze_reason = reason if frozen else ''
        goods.save(update_fields=['is_frozen', 'freeze_reason', 'updated_at'])

        logger.info(f"User {request.user.username} {'froze' if frozen else 'unfroze'} goods {goods.code}")

        return success_response(
            data=GoodsSerializer(goods).data,
            message='冻结成功' if frozen else '解冻成功'
        )


class InvestigationListView(APIView):
    """调查记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = Investigation.objects.all().order_by('-opened_at')

        goods_id = request.query_params.get('goods_id')
        if goods_id:
            queryset = queryset.filter(goods_id=goods_id)
        status = request.query_params.get('status')
        if status:
            queryset = queryset.filter(status=status)

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        serializer = InvestigationSerializer(queryset[start:end], many=True)

        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })

    def post(self, request):
        """调查立案"""
        serializer = InvestigationCreateSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error(serializer.errors) or '参数校验失败')

        investigation = Investigation.objects.create(
            goods_id=serializer.validated_data['goods'],
            title=serializer.validated_data['title'],
            remark=serializer.validated_data['remark'],
            opened_by=request.user
        )

        logger.info(f"User {request.user.username} opened investigation {investigation.id} on goods {investigation.goods_id}")

        return success_response(data=InvestigationSerializer(investigation).data, message='立案成功')


class InvestigationCloseView(APIView):
    """调查结案视图"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            investigation = Investigation.objects.get(pk=pk)
        except Investigation.DoesNotExist:
            return error_response(message='调查记录不存在', code=404)

        if investigation.status == 'closed':
            return error_response(message='调查已结案')

        investigation.status = 'closed'
        investigation.closed_by = request.user
        investigation.closed_at = timezone.now()
        investigation.save(update_fields=['status', 'closed_by', 'closed_at'])

        logger.info(f"User {request.user.username} closed investigation {investigation.id}")

        return success_response(data=InvestigationSerializer(investigation).data, message='结案成功')


class DestructionPlanListView(APIView):
    """销毁计划列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = DestructionPlan.objects.all().order_by('-created_at')

        status = request.query_params.get('status')
        if status:
            queryset = queryset.filter(status=status)

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        serializer = DestructionPlanSerializer(queryset[start:end], many=True)

        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })

    def post(self, request):
        """创建销毁计划（第一阶段：销毁计划）"""
        serializer = DestructionPlanCreateSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error(serializer.errors) or '参数校验失败')

        with transaction.atomic():
            plan = DestructionPlan.objects.create(
                plan_no=uuid.uuid4().hex[:30],
                title=serializer.validated_data['title'],
                remark=serializer.validated_data['remark'],
                created_by=request.user
            )
            plan.plan_no = f"DEST-{timezone.now():%Y%m%d}-{plan.pk:04d}"
            plan.save(update_fields=['plan_no'])
            DestructionPlanItem.objects.bulk_create([
                DestructionPlanItem(
                    plan=plan,
                    goods_id=item['goods'],
                    quantity=item['quantity'],
                    reason=item['reason'],
                    method=item['method'],
                    remark=item['remark']
                )
                for item in serializer.validated_data['items']
            ])
            _save_plan_snapshot(plan, 'plan', request.user)

        logger.info(f"User {request.user.username} created destruction plan {plan.plan_no}")

        return success_response(data=DestructionPlanSerializer(plan).data, message='创建成功')


class DestructionPlanDetailView(APIView):
    """销毁计划详情视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            plan = DestructionPlan.objects.get(pk=pk)
        except DestructionPlan.DoesNotExist:
            return error_response(message='销毁计划不存在', code=404)

        return success_response(data={
            'plan': DestructionPlanSerializer(plan).data,
            'snapshots': DestructionSnapshotSerializer(plan.snapshots.all(), many=True).data,
            'corrections': DestructionCorrectionSerializer(plan.corrections.all(), many=True).data,
        })


class DestructionPlanReviewView(APIView):
    """销毁资格复核视图（第二阶段：资格复核）"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            plan = DestructionPlan.objects.get(pk=pk)
        except DestructionPlan.DoesNotExist:
            return error_response(message='销毁计划不存在', code=404)

        if plan.status != 'planning':
            return error_response(message=f'当前状态（{plan.get_status_display()}）不允许复核')

        remark = str(request.data.get('remark', '')).strip()
        results = []

        with transaction.atomic():
            items = plan.items.select_related('goods').filter(status='pending').order_by('id')
            for item in items:
                problems = _destruction_blockers(item.goods, item)
                if problems:
                    item.status = 'ineligible'
                    item.review_note = '；'.join(problems)
                else:
                    item.status = 'eligible'
                    item.review_note = '复核通过'
                item.save(update_fields=['status', 'review_note', 'updated_at'])
                results.append({
                    'item': item.id,
                    'goods_code': item.goods.code,
                    'status': item.status,
                    'note': item.review_note
                })

            plan.status = 'reviewed'
            plan.reviewed_by = request.user
            plan.reviewed_at = timezone.now()
            plan.review_remark = remark
            plan.save(update_fields=['status', 'reviewed_by', 'reviewed_at', 'review_remark', 'updated_at'])
            _save_plan_snapshot(plan, 'review', request.user)

        logger.info(f"User {request.user.username} reviewed destruction plan {plan.plan_no}")

        return success_response(data={
            'plan': DestructionPlanSerializer(plan).data,
            'results': results
        }, message='复核完成')


class DestructionPlanApproveView(APIView):
    """销毁批准视图（第三阶段：批准）"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            plan = DestructionPlan.objects.get(pk=pk)
        except DestructionPlan.DoesNotExist:
            return error_response(message='销毁计划不存在', code=404)

        if plan.status != 'reviewed':
            return error_response(message=f'当前状态（{plan.get_status_display()}）不允许批准')

        decision = request.data.get('decision')
        if decision not in ('approved', 'rejected'):
            return error_response(message='请指定批准结论（decision: approved/rejected）')
        remark = str(request.data.get('remark', '')).strip()

        if decision == 'approved' and not plan.items.filter(status='eligible').exists():
            return error_response(message='没有复核通过的明细，无法批准')

        with transaction.atomic():
            plan.status = decision
            plan.approved_by = request.user
            plan.approved_at = timezone.now()
            plan.approve_remark = remark
            plan.save(update_fields=['status', 'approved_by', 'approved_at', 'approve_remark', 'updated_at'])
            _save_plan_snapshot(plan, 'approve', request.user)

        logger.info(f"User {request.user.username} {decision} destruction plan {plan.plan_no}")

        return success_response(
            data=DestructionPlanSerializer(plan).data,
            message='批准成功' if decision == 'approved' else '已驳回'
        )


class DestructionPlanExecuteView(APIView):
    """销毁执行确认视图（第四阶段：执行确认，执行后进入不可逆终态）"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        try:
            plan = DestructionPlan.objects.get(pk=pk)
        except DestructionPlan.DoesNotExist:
            return error_response(message='销毁计划不存在', code=404)

        if plan.status != 'approved':
            return error_response(message=f'当前状态（{plan.get_status_display()}）不允许执行')

        results = []
        # 每条明细在独立事务中处理：单项失败不会波及其他明细，每项都有明确结果
        for item in plan.items.filter(status='eligible').order_by('id'):
            try:
                with transaction.atomic():
                    locked = DestructionPlanItem.objects.select_for_update().get(pk=item.pk)
                    if locked.status != 'eligible':
                        results.append({
                            'item': locked.id,
                            'status': locked.status,
                            'note': '明细状态已变化，跳过执行'
                        })
                        continue
                    goods = Goods.objects.select_for_update().get(pk=locked.goods_id)
                    # 执行前再次核查冻结、未结调查与领用状态
                    problems = _destruction_blockers(goods, locked)
                    if problems:
                        locked.status = 'blocked'
                        locked.execute_note = '；'.join(problems)
                    else:
                        goods.quantity -= locked.quantity
                        goods.save(update_fields=['quantity', 'updated_at'])
                        locked.status = 'destroyed'
                        locked.destroyed_at = timezone.now()
                        locked.execute_note = '执行完成'
                    locked.save(update_fields=['status', 'execute_note', 'destroyed_at', 'updated_at'])
                    results.append({
                        'item': locked.id,
                        'goods_code': goods.code,
                        'status': locked.status,
                        'note': locked.execute_note
                    })
            except Exception as exc:
                logger.exception(f"Destruction plan {plan.plan_no} item {item.id} execution failed")
                # 在独立事务中将该项标记为失败，不影响其他明细的处理结果
                with transaction.atomic():
                    locked = DestructionPlanItem.objects.get(pk=item.pk)
                    locked.status = 'failed'
                    locked.execute_note = f'处理异常：{exc}'[:300]
                    locked.save(update_fields=['status', 'execute_note', 'updated_at'])
                results.append({
                    'item': locked.id,
                    'status': 'failed',
                    'note': locked.execute_note
                })

        with transaction.atomic():
            plan.status = 'executed'
            plan.executed_by = request.user
            plan.executed_at = timezone.now()
            plan.save(update_fields=['status', 'executed_by', 'executed_at', 'updated_at'])
            _save_plan_snapshot(plan, 'execute', request.user)

        destroyed_count = sum(1 for r in results if r['status'] == 'destroyed')
        blocked_count = sum(1 for r in results if r['status'] == 'blocked')
        failed_count = sum(1 for r in results if r['status'] == 'failed')

        logger.info(
            f"User {request.user.username} executed destruction plan {plan.plan_no}: "
            f"destroyed={destroyed_count} blocked={blocked_count} failed={failed_count}"
        )

        return success_response(data={
            'plan': DestructionPlanSerializer(plan).data,
            'results': results,
            'destroyed_count': destroyed_count,
            'blocked_count': blocked_count,
            'failed_count': failed_count
        }, message=f'执行完成：销毁 {destroyed_count} 项，拦截 {blocked_count} 项，失败 {failed_count} 项')


class DestructionCorrectionListView(APIView):
    """销毁更正视图（仅修改元数据，不恢复库存）"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            plan = DestructionPlan.objects.get(pk=pk)
        except DestructionPlan.DoesNotExist:
            return error_response(message='销毁计划不存在', code=404)

        serializer = DestructionCorrectionSerializer(plan.corrections.all(), many=True)
        return success_response(data=serializer.data)

    def post(self, request, pk):
        try:
            plan = DestructionPlan.objects.get(pk=pk)
        except DestructionPlan.DoesNotExist:
            return error_response(message='销毁计划不存在', code=404)

        if plan.status != 'executed':
            return error_response(message='仅已执行的销毁计划可以提交更正')

        serializer = DestructionCorrectionCreateSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error(serializer.errors) or '参数校验失败')

        try:
            item = plan.items.get(pk=serializer.validated_data['item'])
        except DestructionPlanItem.DoesNotExist:
            return error_response(message='销毁明细不存在', code=404)
        if item.status != 'destroyed':
            return error_response(message='仅已销毁的明细可以更正')

        changes = serializer.validated_data['changes']
        with transaction.atomic():
            before_after = {}
            for field, new_value in changes.items():
                before_after[field] = {'old': getattr(item, field), 'new': new_value}
                setattr(item, field, new_value)
            item.save(update_fields=list(changes.keys()) + ['updated_at'])
            correction = DestructionCorrection.objects.create(
                plan=plan,
                item=item,
                changes=before_after,
                reason=serializer.validated_data['reason'],
                created_by=request.user
            )

        logger.info(
            f"User {request.user.username} corrected destruction plan {plan.plan_no} "
            f"item {item.id}: {list(changes.keys())}"
        )

        return success_response(data=DestructionCorrectionSerializer(correction).data, message='更正成功')
